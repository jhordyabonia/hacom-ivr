from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SOURCE = Path('/media/jhordy/documents/hacom/ivr/results/cruce_RFP_maqueta_FreeSWITCH.md')
OUTPUT = SOURCE.with_suffix('.pptx')

NAVY = RGBColor(22, 58, 90)
BLUE = RGBColor(32, 120, 180)
TEAL = RGBColor(26, 156, 147)
GOLD = RGBColor(214, 167, 59)
LIGHT = RGBColor(245, 247, 250)
WHITE = RGBColor(255, 255, 255)
TEXT = RGBColor(40, 48, 58)
GRAY = RGBColor(110, 120, 130)
GREEN = RGBColor(34, 122, 84)
RED = RGBColor(148, 52, 44)


def add_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_top_bar(slide, color=NAVY):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 13.333, Inches(0.42))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()


def add_title(slide, text, left=0.6, top=0.55, width=11.0, height=0.5, size=24, color=NAVY):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color


def add_subtitle(slide, text, left=0.6, top=1.08):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(11.5), Inches(0.35))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    run = p.runs[0]
    run.font.size = Pt(11)
    run.font.italic = True
    run.font.color.rgb = GRAY


def add_footer(slide, text='uVAS 2026 | FreeSWITCH + IMS | HACOM'):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12), Inches(0.25))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.RIGHT
    run = p.runs[0]
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY


def add_bullets(slide, lines, left, top, width, height, font_size=16, bullet=True, color=TEXT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        if bullet:
            p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        run = p.runs[0]
        run.font.size = Pt(font_size)
        run.font.color.rgb = color


def add_callout(slide, title, text, left, top, width, height, fill_color=LIGHT, border_color=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.2)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    run = p.runs[0]
    run.font.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = NAVY
    p2 = tf.add_paragraph()
    p2.text = text
    run2 = p2.runs[0]
    run2.font.size = Pt(11)
    run2.font.color.rgb = TEXT


def add_metric_card(slide, title, value, subtitle, left, top, width, height, accent=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = accent
    shape.line.width = Pt(1.2)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    run = p.runs[0]
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = GRAY
    p2 = tf.add_paragraph()
    p2.text = value
    run2 = p2.runs[0]
    run2.font.size = Pt(22)
    run2.font.bold = True
    run2.font.color.rgb = accent
    p3 = tf.add_paragraph()
    p3.text = subtitle
    run3 = p3.runs[0]
    run3.font.size = Pt(9)
    run3.font.color.rgb = TEXT


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Slide 1: portada
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)

badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(0.7), Inches(2.4), Inches(0.7))
badge.fill.solid(); badge.fill.fore_color.rgb = NAVY; badge.line.fill.background()
badge_tf = badge.text_frame
p = badge_tf.paragraphs[0]; p.text = 'HACOM'; p.alignment = PP_ALIGN.CENTER; run = p.runs[0]; run.font.size = Pt(26); run.font.bold = True; run.font.color.rgb = WHITE

add_title(slide, 'Cruce RFP uVAS 2026', left=0.8, top=1.6, width=7.5, size=28)
add_title(slide, '↔ Maqueta FreeSWITCH', left=0.8, top=2.1, width=7.5, size=26)

box = slide.shapes.add_textbox(Inches(0.8), Inches(2.9), Inches(7.0), Inches(2.2))
tf = box.text_frame; tf.word_wrap = True
for i, text in enumerate([
    'Documento base: cruce_RFP_maqueta_FreeSWITCH.md',
    'Objetivo: validar qué funcionalidades del RFP están cubiertas por la maqueta actual.',
    'Conclusión: la plataforma cubre hoy voz (IVR/VMS) y prepara el resto de servicios del RFP con la ruta de mensajería y AS complementarios.',
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = text
    run = p.runs[0]; run.font.size = Pt(18 if i == 0 else 16); run.font.color.rgb = TEXT

panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.8), Inches(1.3), Inches(3.7), Inches(4.0))
panel.fill.solid(); panel.fill.fore_color.rgb = NAVY; panel.line.fill.background()
ptf = panel.text_frame; ptf.word_wrap = True
for i, line in enumerate(['RFP', 'uVAS 2026', 'Core IMS', 'Kamailio', 'FreeSWITCH AS', 'RTPEngine', 'Operación via mims_ui']):
    p = ptf.paragraphs[0] if i == 0 else ptf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = line
    run = p.runs[0]
    run.font.size = Pt(18 if i in (0, 2, 4, 6) else 16)
    run.font.bold = i in (0, 2, 4, 6)
    run.font.color.rgb = WHITE
add_footer(slide)

# Slide 2: Resumen ejecutivo
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_title(slide, 'Resumen ejecutivo')
add_subtitle(slide, 'La maqueta ya valida la base funcional de voz y ofrece la plataforma para cubrir el resto del RFP')

add_metric_card(slide, 'Cobertura funcional hoy', '2/12', 'IVR y VMS validados', left=0.7, top=1.6, width=2.2, height=1.8, accent=GREEN)
add_metric_card(slide, 'Cobertura nativa ampliable', '6/12', 'SCE, MCA, charging, reporting, monitoring, cloud-ready', left=3.3, top=1.6, width=3.2, height=1.8, accent=BLUE)
add_metric_card(slide, 'Requiere AS complementario', '4/12', 'SMSC, FDA, SMSFW, USSD', left=7.0, top=1.6, width=2.8, height=1.8, accent=GOLD)
add_metric_card(slide, 'Verificación E2E', '15/15', 'pass con scripts de validación', left=10.1, top=1.6, width=2.4, height=1.8, accent=TEAL)

add_bullets(slide, [
    'El RFP solicita una plataforma uVAS convergente con SMSC, FDA, SMSFW, VMS, MCA, USSD, IVR, SCE, charging, reporting y monitoring.',
    'La maqueta incluye Core IMS con Kamailio + PyHSS + AS Frontera + FreeSWITCH como AS aplicación + RTPEngine + mims_ui.',
    'La conclusión adelantada es que la maqueta cubre hoy las funcionalidades de voz y prepara el resto del RFP ampliando con capacidad nativa de FreeSWITCH o AS complementarios dentro de la misma arquitectura por iFC.',
    'La clave técnica es que ningún cambio crítico en el Core IMS es necesario: el enrutado de servicios se hace a través de iFC y AS Frontera.',
], left=0.8, top=3.9, width=11.8, height=2.5, font_size=15)
add_footer(slide)

# Slide 3: Funcionalidades esperadas por el RFP
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_title(slide, '1. Funcionalidades esperadas por el RFP')
add_subtitle(slide, 'Sección 1.1 Background, página 24')

lines = [
    'F1 — SMSC (Short Message Service Center) | Mensajería',
    'F2 — FDA — First Delivery Attempt | Mensajería',
    'F3 — SMS Firewall (SMSFW) | Mensajería / Seguridad',
    'F4 — Voice Mail System (VMS) | Voz',
    'F5 — Missed Call Alert (MCA) | Voz / Notificación',
    'F6 — USSD Gateway | Servicios interactivos',
    'F7 — Interactive Voice Response (IVR) | Voz',
    'F8 — SCE — Service Creation Environment | Desarrollo de servicios',
    'F9 — Charging (tarificación/cobro) | Horizontal',
    'F10 — Reporting (informes) | Horizontal',
    'F11 — Monitoring & Management | Horizontal',
    'F12 — Convergencia / cloud-native / virtualizado / 5G-ready | Plataforma',
]
add_bullets(slide, lines, 0.8, 1.6, 11.8, 5.2, font_size=15)
add_footer(slide)

# Slide 4: Qué es la maqueta FreeSWITCH
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_title(slide, '2. La maqueta de FreeSWITCH (qué es y qué está validado)')

add_bullets(slide, [
    'Core IMS: Kamailio P-CSCF (:5060), I-CSCF (:4060), S-CSCF (:6060) con Cx Diameter vía PyHSS (UAR/UAA, MAR/MAA, SAR/SAA + iFC).',
    'AS Frontera: Kamailio asfront, interworking ISC → SIP plano, destino del iFC#1.',
    'AS aplicación: FreeSWITCH (mims3_freeswitch) con dialplan public para IVR y VMS nativos, sin Asterisk.',
    'Plano de medios: RTPEngine anclado en el P-CSCF; autenticación Digest-AKAv1-MD5 E2E demostrada; registro IMS 401 → 200.',
    'Orquestación: docker-compose con red 172.32.0.0/24, 14 contenedores, restart unless-stopped.',
    'Interfaz de operación: mims_ui en http://localhost:8090 para estado, E2E, operación, logs y VMS.',
    'Verificación: scripts/test_maqueta.sh → PASS 15 / FAIL 0 (registro, iFC, IVR, RTPEngine, VMS).',
    'Evidencia VMS: freeswitch/vms/0100003/INBOX/msg_*.wav; evidencia IVR: playback de ivr_bienvenida.wav a 0100002; evidencia AKA: 401 con algorithm=AKAv1-MD5 sin ck/ik → 200 OK.',
], left=0.8, top=1.5, width=12.0, height=5.3, font_size=14)
add_footer(slide)

# Slide 5: Matriz de cruce RFP ↔ maqueta FreeSWITCH (summary)
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_title(slide, '3. Matriz de cruce RFP ↔ maqueta FreeSWITCH')
add_subtitle(slide, 'Estado: ✅ Implementado • 🟡 Ampliable-nativo • 🟠 Complemento-AS')

matrix_lines = [
    'F7 IVR | ✅ Implementado | mod_dptools + mod_say_* + mod_lua + mod_httapi + mod_xml_curl | Evidencia: public.xml => ivr_0100002',
    'F4 VMS | ✅ Implementado | mod_voicemail + record + persistencia en storage/vms | Evidencia: msg_*.wav en freeswitch/vms/0100003',
    'F8 SCE | 🟡 Ampliable-nativo | Lua, httapi, xml_curl, flujos JSON por país | Módulos cargados: mod_lua, mod_httapi',
    'F5 MCA | 🟡 Ampliable-nativo | CDR + no-answer/busy detection + SMS/MB/MWI | mod_cdr_csv + mod_event_socket + mod_sms',
    'F1 SMSC | 🟠 Complemento-AS | mensaje MESSAGE por iFC + AS SMS o mod_sms | mismo patrón ISC validado para INVITE',
    'F2 FDA | 🟠 Complemento-AS | cola persistida y reintento diferido | redis/MySQL en la maqueta',
    'F3 SMS Firewall | 🟠 Complemento-AS | MO/MT filters, rate limit, blacklist, scoring | AS smsfw interpuesto en la cadena',
    'F6 USSD Gateway | 🟠 Complemento-AS | menú por etapas sobre SIP MESSAGE | motor IVR/SCE reutilizado',
    'F9 Charging | 🟡 Ampliable | mod_cdr_csv, mod_xml_cdr, mod_json_cdr | API de ingesta a BD/CDR',
    'F10 Reporting | 🟡 Ampliable | mims_ui consultando CDR/MySQL + mod_db | interfaz 8090',
    'F11 Monitor & Management | 🟡 Ampliable | mod_event_socket + mims_ui + healthchecks | docker compose / docker ps',
    'F12 Convergencia 5G-ready | 🟡 Ampliable | contenerización + 14 servicios + camino a K8s/Helm | arquitectura preparada',
]
add_bullets(slide, matrix_lines, 0.8, 1.5, 12.0, 5.2, font_size=12)
add_footer(slide)

# Slide 6: IVR y VMS detalles
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_title(slide, '4. Detalle por funcionalidad: IVR y VMS')

add_callout(slide, 'F7 — IVR ✅', 'public.xml: ivr_0100002 → answer → sleep → playback(ivr_bienvenida.wav) → hangup. El menú DTMF se construye con mod_dptools, mod_lua/mod_httapi consultando MySQL y mod_say_* para lectura del saldo o número. RTP anclado en RTPEngine.', 0.7, 1.5, 5.8, 2.2)
add_callout(slide, 'F4 — VMS ✅', 'public.xml: vms_0100003 → graba el mensaje RTP del llamante a vms/<MSISDN>/INBOX/msg_<fecha>.wav. mod_voicemail cargado da buzón completo, greeting y MWI/message-summary. Volumen persistente ./freeswitch/vms.', 0.7, 4.0, 5.8, 2.5)
add_bullets(slide, [
    'La ruta del flujo es UE → P-CSCF → S-CSCF → iFC → AS Frontera → FreeSWITCH.',
    'Se valida la parte de voz con RTPEngine y la autenticación AKAv1-MD5 E2E.',
    'Cobertura TS 26.114: mod_amr/mod_amrwb cargados; el softphone actual no ofrece AMR, pero no es limitación de la maqueta.',
], 7.1, 1.8, 5.3, 4.0, font_size=14)
add_footer(slide)

# Slide 7: SMS / FDA / SMSFW / USSD / MCA
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_title(slide, '5. Dominio de mensajería y servicios interactivos')

add_bullets(slide, [
    'F5 — MCA 🟡: FreeSWITCH genera CDR con mod_cdr_csv; detecta disposition=no answer / busy. El aviso se materializa por SMS con mod_sms o por MB/MWI vía S-CSCF.',
    'F1 — SMSC 🟠: el mismo mecanismo ISC que enruta INVITE puede enrutar MESSAGE por iFC con Method=MESSAGE y prioridad propia en maqueta_ifc.xml. El AS SMS normaliza, persiste y entrega; stubs SMPP para integrar el SMSC real del operador.',
    'F2 — FDA 🟠: entrega inmediata si el destino está registrado; diferida si no lo está, reintenta con el siguiente REGISTER. Estados delivered/undelivered/deleted en BD.',
    'F3 — SMS Firewall 🟠: interpuesto en la cadena del SMSC; filtros MO/MT por número, contenido, blacklist, rate-limit y scoring. Rechazo o reescritura antes de entregar.',
    'F6 — USSD Gateway 🟠: reutiliza el motor de menú del IVR sobre SIP MESSAGE con persistencia de sesión (redis/MySQL). El enlace real al USSDC SIGTRAN/MAP queda como integración de propuesta.',
], left=0.8, top=1.5, width=12.0, height=5.4, font_size=14)
add_footer(slide)

# Slide 8: SCE, Charging, Reporting, Monitoring, Cloud
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_title(slide, '6. Capacidad nativa de FreeSWITCH y control del servicio')

add_bullets(slide, [
    'F8 — SCE 🟡: FreeSWITCH ofrece motor nativo con mod_lua, mod_httapi, mod_xml_curl y flujos JSON/YAML/Lua; servicios desplegables sin redeploy del Core IMS y sin tocar la lógica del S-CSCF.',
    'F9 — Charging 🟡: mod_cdr_csv cargado; mod_xml_cdr/mod_json_cdr en disco permiten emitir CDR a la API de facturación. Persistencia en MySQL. Ro/Rf como fases posteriores.',
    'F10 — Reporting 🟡: mims_ui (puerto 8090) ya consulta estado de servicio, CDR y MySQL. El entorno soporta tablas por rellenar con CDR del AS.',
    'F11 — Monitoring & Management 🟡: mod_event_socket expone eventos y estadísticas; mims_ui supervisa contenedores, DNS, listeners, E2E, operación y VMS. Healthchecks por compose.',
    'F12 — Convergencia / cloud-native / 5G-ready 🟡: todo está contenerizado por servicio, con red parametrizada, 14 contenedores y capacidad de escalar vía Docker/Kubernetes y fases de SBA/NEF.',
], left=0.8, top=1.5, width=12.0, height=5.4, font_size=14)
add_footer(slide)

# Slide 9: módulos FreeSWITCH verificados
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_title(slide, '7. Módulos FreeSWITCH de la maqueta (verificado en el contenedor)')

add_bullets(slide, [
    'Cargados: mod_sofia, mod_dptools, mod_voicemail, mod_amr, mod_amrwb, mod_opus, mod_cdr_csv, mod_conference, mod_fifo, mod_lua, mod_httapi, mod_loopback, mod_event_socket, mod_dialplan_xml, mod_say_*, mod_local_stream, mod_sndfile, mod_tone_stream, mod_spandsp.',
    'En disco / habilitables: mod_sms.so, mod_xml_curl.so, mod_json_cdr.so, mod_cdr_mongodb.so, mod_java.so, mod_perl.so y otros módulos de integración.',
    'Nota crucial del diseño: no existe mod_smpp.so obligado para el patrón de mensajería; el tránsito SIP MESSAGE → FreeSWITCH/mod_sms o AS python cubre la misma función sin depender de ese módulo.',
    'El patrón por iFC se mantiene constante para voz y mensajería, lo cual reduce la superficie de cambio del Core IMS y acelera el desarrollo.',
], left=0.8, top=1.5, width=12.0, height=5.3, font_size=14)
add_footer(slide)

# Slide 10: conclusiones y propuesta técnica
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_title(slide, '8. Conclusión')

add_bullets(slide, [
    '1. Voz del AS — cumplido hoy: IVR y VMS funcionando y validados E2E; el flujo está enrutado por iFC del S-CSCF hacia FreeSWITCH y con medios anclados en RTPEngine.',
    '2. Plataforma convergente — la maqueta es la base para el resto de funcionalidades del RFP sin cambiar el Core IMS: SCE, MCA, charging, reporting y monitoring se apoyan en módulos y servicios nativos de FreeSWITCH.',
    '3. Dominio de mensajería — SMSC, FDA, SMSFW y USSD se resuelven sobre el mismo patrón iFC validado (MESSAGE) junto con un AS de mensajería dentro de la misma maqueta.',
    '4. Ventaja competitiva — un solo plano VAS para voz y servicios, una ruta de entrega más corta y una base técnica sólida para la propuesta del operador.',
    '5. La evidencia demuestra que la maqueta hoy no sólo es un prototipo; es una base operativa para llegar a una propuesta uVAS realista, convergente y 5G-ready.',
], left=0.8, top=1.5, width=12.0, height=5.5, font_size=14)
add_footer(slide)

# Slide 11: referencias
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)
add_title(slide, '9. Referencias y evidencia')

refs = [
    'RFP: doc/20260817 uVAS RFP 2026.md — funcionalidades en §1.1, página 24.',
    'Plan de etapas: results/plan_trabajo_uVAS_2026.md — etapas 1 a 6; F1-F4 en etapa 2, USSD en etapa 3, SCE/charging en etapas 1/4/5.',
    'Comparativa AS: results/argumentacion_AS_kamailio_vs_freeswitch.md.',
    'Cumplimiento 3GPP: results/diagnostico_cumplimiento_3gpp.md.',
    'E2E: maqueta-ims-3/scripts/test_maqueta.sh — PASS 15/15.',
    'AKA: maqueta-ims-3/scripts/ue_aka_probe.py — Digest-AKAv1-MD5 → 200 OK.',
    'AS Frontera ISC: maqueta-ims-3/asfront/{asfront.cfg,kamailio_asfront.cfg}.',
    'FreeSWITCH AS: maqueta-ims-3/freeswitch/conf/{dialplan/public.xml, sip_profiles/external.xml}.',
]
add_bullets(slide, refs, 0.8, 1.5, 11.8, 5.5, font_size=14)
add_footer(slide)

prs.save(OUTPUT)
print(f'Presentación creada en {OUTPUT}')
