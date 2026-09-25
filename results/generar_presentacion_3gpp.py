from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SOURCE = Path('/media/jhordy/documents/hacom/ivr/results/diagnostico_cumplimiento_3gpp.md')
OUTPUT = SOURCE.with_suffix('.pptx')

NAVY = RGBColor(17, 50, 86)
BLUE = RGBColor(32, 121, 181)
TEAL = RGBColor(24, 143, 142)
GREEN = RGBColor(40, 128, 76)
GOLD = RGBColor(212, 172, 60)
RED = RGBColor(167, 57, 44)
LIGHT = RGBColor(244, 247, 249)
WHITE = RGBColor(255, 255, 255)
TEXT = RGBColor(35, 44, 58)
GRAY = RGBColor(107, 118, 130)


def add_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_top_bar(slide, color=NAVY):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 13.333, Inches(0.42))
    bar.fill.solid(); bar.fill.fore_color.rgb = color
    bar.line.fill.background()


def add_title(slide, text, left=0.6, top=0.55, width=11.0, height=0.5, size=24, color=NAVY):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color


def add_subtitle(slide, text, left=0.6, top=1.08):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(11.5), Inches(0.35))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    run = p.runs[0]
    run.font.size = Pt(11)
    run.font.italic = True
    run.font.color.rgb = GRAY


def add_footer(slide, text='IMS / 3GPP compliance | HACOM 2026'):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12), Inches(0.25))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.RIGHT
    run = p.runs[0]
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY


def add_bullets(slide, lines, left, top, width, height, font_size=15, color=TEXT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        run = p.runs[0]
        run.font.size = Pt(font_size)
        run.font.color.rgb = color


def add_callout(slide, title, text, left, top, width, height, fill=LIGHT, border=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border; shape.line.width = Pt(1.2)
    tf = shape.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    run = p.runs[0]; run.font.bold = True; run.font.size = Pt(15); run.font.color.rgb = NAVY
    p2 = tf.add_paragraph(); p2.text = text; run2 = p2.runs[0]; run2.font.size = Pt(11); run2.font.color.rgb = TEXT


def add_metric_card(slide, title, value, subtitle, left, top, width, height, accent=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid(); shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = accent; shape.line.width = Pt(1.2)
    tf = shape.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title
    run = p.runs[0]; run.font.size = Pt(11); run.font.bold = True; run.font.color.rgb = GRAY
    p2 = tf.add_paragraph(); p2.text = value
    run2 = p2.runs[0]; run2.font.size = Pt(22); run2.font.bold = True; run2.font.color.rgb = accent
    p3 = tf.add_paragraph(); p3.text = subtitle
    run3 = p3.runs[0]; run3.font.size = Pt(9); run3.font.color.rgb = TEXT


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# portada
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_top_bar(slide)

badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(0.7), Inches(2.8), Inches(0.7))
badge.fill.solid(); badge.fill.fore_color.rgb = NAVY; badge.line.fill.background()
pt = badge.text_frame.paragraphs[0]; pt.text = 'HACOM'; pt.alignment = PP_ALIGN.CENTER; run = pt.runs[0]; run.font.bold = True; run.font.size = Pt(26); run.font.color.rgb = WHITE

add_title(slide, 'Diagnóstico de cumplimiento 3GPP', left=0.8, top=1.6, width=7.5, size=27)
add_title(slide, 'maqueta-ims-3 — estado tras corrección', left=0.8, top=2.2, width=8.1, size=24)

box = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(7.2), Inches(2.2))
tf = box.text_frame; tf.word_wrap = True
for i, text in enumerate([
    'Documento vivo: diagnóstico original del tráfico SIP/IMS y estado de corrección.',
    'Objetivo: validar los puntos de fallo reportados y dejar evidencia reproducible del estado actual.',
    'Resultado: la maqueta corrige los incumplimientos críticos de autenticación y registro IMS, quedando en un estado útil para pruebas de validación 3GPP.',
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = text; run = p.runs[0]; run.font.size = Pt(18 if i == 0 else 16); run.font.color.rgb = TEXT

panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.0), Inches(1.4), Inches(3.6), Inches(4.2))
panel.fill.solid(); panel.fill.fore_color.rgb = NAVY; panel.line.fill.background()
panel_tf = panel.text_frame; panel_tf.word_wrap = True
for i, line in enumerate(['3GPP', 'TS 33.203', 'TS 24.229', 'TS 23.003', 'SIP / IMS', 'E2E 15/15']):
    p = panel_tf.paragraphs[0] if i == 0 else panel_tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER; p.text = line
    run = p.runs[0]; run.font.bold = i in (0, 2, 4); run.font.size = Pt(18 if i in (0, 4) else 16); run.font.color.rgb = WHITE
add_footer(slide)

# slide 2 summary metrics
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide); add_top_bar(slide)
add_title(slide, 'Resumen ejecutivo')
add_subtitle(slide, 'Actualización tras la corrección de los incumplimientos críticos')

add_metric_card(slide, 'Autenticación AKA', 'CORREGIDO', 'AKAv1-MD5 validado por sonda', left=0.7, top=1.6, width=2.6, height=1.8, accent=GREEN)
add_metric_card(slide, 'Seguridad Gm / SIP', 'PARCIAL', 'Nivel SIP OK; túnel ESP pendiente', left=3.7, top=1.6, width=2.8, height=1.8, accent=GOLD)
add_metric_card(slide, 'Contact / identity IMS', 'CORREGIDO', 'IMPU normalizado + feature-tags', left=7.0, top=1.6, width=2.8, height=1.8, accent=BLUE)
add_metric_card(slide, 'E2E validado', '15/15', 'test_maqueta.sh PASS', left=10.2, top=1.6, width=2.2, height=1.8, accent=TEAL)

add_bullets(slide, [
    'El principal problema crítico corregido fue la autenticación: el 401 original exigía MD5 y la corrección lleva el flujo a AKAv1-MD5 con la identidad adecuada.',
    'La negociación SIP-level de seguridad ya se materializa con Security-Client / Server / Verify; el túnel ESP real sigue pendiente por limitación del entorno Docker.',
    'Los registros IMS usan IMPU normalizado (MSISDN) y no IMSI, cumpliendo la práctica 3GPP y evitando el 403 de HSS User Unknown.',
    'La validación final con pruebas E2E confirma PASS: 15 / FAIL: 0.',
], left=0.8, top=3.9, width=11.7, height=2.4, font_size=15)
add_footer(slide)

# slide 3 critical issues
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide); add_top_bar(slide)
add_title(slide, '1. Incumplimientos críticos — Estado tras corrección')
add_subtitle(slide, 'Puntos de fallo reportados y resolución aplicada')

add_callout(slide, '1.1 Autenticación AKA', 'Problema original: 401 exigía algorithm=MD5. Violación TS 33.203. Corrección: S-CSCF aplica AKAv1-MD5; PyHSS y UAR/MAR normalizan por MSISDN; la sonda completa Digest-AKAv1-MD5 extremo a extremo con 200 OK.', 0.7, 1.5, 5.7, 2.2)
add_callout(slide, '1.2 Seguridad Gm / IPSec', 'La negociación SIP-level ya ocurre con Security-Client / Security-Server / Security-Verify. El túnel ESP real no se materializa por falta de NET_ADMIN en Docker, así que la capa señalización queda resuelta pero no la capa IPsec real.', 0.7, 4.1, 5.7, 2.2)
add_callout(slide, '1.3 Contact + feature-tags', 'Se corrige el uso de Contact con IP directa y se incorporan atributos IMS como +g.3gpp.icsi-ref y +g.3gpp.smsip. El registro pasa a tener sentido 3GPP en la capa de control.', 6.8, 1.5, 5.8, 2.2)
add_callout(slide, '1.4 Identities SIP / IMPU + IMPI', 'Antes se usaba IMSI en URI. Ahora el registro usa IMPU normalizado sip:0010100001@… y el Authorization conserva la identidad privada (IMPI). El fallo HSS User Unknown queda corregido.', 6.8, 4.1, 5.8, 2.2)
add_footer(slide)

# slide 4 algorithms and evidence
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide); add_top_bar(slide)
add_title(slide, '1.1 Algoritmo de autenticación (TS 33.203)')

add_bullets(slide, [
    'Problema original: 401 Unauthorized exigía algorithm=MD5.',
    'Violación: 3GPP exige AKA-v1-MD5 / AKA-v2-SHA-256 con vector USIM/ISIM.',
    'Corrección aplicada: scscf/kamailio_scscf.cfg autentica con AKAv1-MD5 si viene Security-Client; si no, aplica MD5 como fallback.',
    'scscf/scscf.cfg activa REG_AUTH_DEFAULT_ALG = "AKAv1-MD5" y desactiva HSS-Selected.',
    'La interfaz de provisión crea el vector auc con opc y sqn correctos; PyHSS normaliza UAR/MAR por MSISDN si el username del Cx no es IMSI de 15 dígitos.',
    'Evidencia: scripts/ue_aka_probe.py completa Digest-AKAv1-MD5 extremo a extremo con 200 OK. El flujo recibe WWW-Authenticate con algorithm=AKAv1-MD5, ck/ik, qop="auth".',
    'Limitación real: pjsua 2.14 no completa AKA estándar por usar password=KI en lugar de XRES; la vía idónea para este tipo de demo es la sonda, no el cliente PJSUA.',
], left=0.8, top=1.5, width=12.0, height=5.5, font_size=14)
add_footer(slide)

# slide 5 security details
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide); add_top_bar(slide)
add_title(slide, '1.2 Seguridad de la interfaz Gm / IPSec (TS 33.203)')

add_bullets(slide, [
    'Problema original: ausencia de Security-Client / Security-Server / Security-Verify.',
    'Estado: la negociación SIP-level ya ocurre: UE envía Security-Client: ipsec-3gpp; alg=hmac-sha-1-96; spi…; port…',
    'P-CSCF responde con Security-Server en el 401 y confirma con Security-Verify en el 200 OK (RFC 3329), añadido en pcscf/route/register.cfg.',
    'Ck/ik ya no se filtran al UE: el strip que estaba dentro de WITH_IPSEC se movió fuera; la 401 hacia el UE llega con algorithm=AKAv1-MD5, qop="auth" sin ck=/ik=.',
    'Colateral corregido: en la 401 ya no se filtran las claves del túnel.',
    'Pendiente documentado: el túnel IPsec ESP real no materializa por falta de privilegios de kernel (NET_ADMIN) en Docker. La negociación Gm queda soportada a nivel de señalización, no de plano de datos.',
], left=0.8, top=1.5, width=12.0, height=5.5, font_size=14)
add_footer(slide)

# slide 6 Contact and identities
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide); add_top_bar(slide)
add_title(slide, '1.3 / 1.4 Formato de Contact e identidades SIP')

add_bullets(slide, [
    'Formato Contact: Contact: <sip:…>;+g.3gpp.icsi-ref="urn:urn-7:3gpp-service.ims.icsi.mmtel";+g.3gpp.smsip. El P-CSCF procesa y guarda el contacto con feature-tags.',
    'Antes: Contact con IP directa …:5061;ob, sin características IMS. Ahora se usa el formato correcto y el registro mantiene la semántica de la identidad IMS.',
    'Antes: From/To usaban IMSI como URI: sip:001011234567890@…',
    'Ahora: el registro usa IMPU normalizado sip:0010100001@ims.mnc001.mcc001.3gppnetwork.org (MSISDN) en From/To; la identidad privada IMPI se declara en Authorization.',
    'scripts/test_maqueta.sh migró a --id sip:$MSISDN@… + --username $MSISDN para mantener el flujo correcto.',
    'Caveat PJSUA: pjsua no envía la identidad privada (IMPI) en el primer REGISTER, por lo que en la vía MD5 el digest usa el MSISDN. La vía AKA conforme la demuestra la sonda ue_aka_probe.py.',
], left=0.8, top=1.4, width=12.0, height=5.7, font_size=14)
add_footer(slide)

# slide 7 INVITE header issues
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide); add_top_bar(slide)
add_title(slide, '2. Desviaciones en el establecimiento de llamada (INVITE)')

add_bullets(slide, [
    '2.1 Cabeceras P-Header (TS 24.229): antes el INVITE saliente no llevaba P-Preferred-Identity ni P-Access-Network-Info.',
    'Corrección aplicada en pcscf/kamailio_pcscf.cfg ruta MO: si el UE no las envía, el P-CSCF inserta P-Access-Network-Info: 3GPP-E-UTRAN-FDD; utran-cell-id-3gpp=0020100000f101 y P-Preferred-Identity: <sip:$fu@dominio>.',
    'También se inserta PANI en el REGISTER en pcscf/route/register.cfg.',
    '2.2 Parámetros SDP / codecs (TS 26.114): el problema original era que SDP no llevaba AMR/AMR-WB.',
    'Estado: la red sí puede servir AMR; FreeSWITCH carga mod_amr.so/mod_amrwb.so y RTPEngine/FS permiten transcodificar a/de AMR.',
    'Limitación del cliente: pjsua de la maqueta se compiló sin códec AMR (solo iLBC/GSM/G722); por ese motivo el SDP de tests E2E continúa sin AMR por parte del UE.',
    'Se requiere un cliente MTSI (Linphone/UCRaw con AMR) para probar AMR extremo a extremo.',
], left=0.8, top=1.5, width=12.0, height=5.5, font_size=14)
add_footer(slide)

# slide 8 summary table-esque text
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide); add_top_bar(slide)
add_title(slide, 'Resumen de evaluación actualizado')

rows = [
    'Flujo Registrar (P-I-S CSCF) | Cumple | Cumple | test 15/15',
    'Cabeceras Path / Service-Route | Cumple | Cumple | test 15/15',
    'Autenticación AKA (AKAv1-MD5) | No cumple | Cumple | sonda UE AKA → 200 OK',
    'Algoritmo 401 por defecto | MD5 | AKAv1-MD5 (fallback MD5) | 401 capturada',
    'Security-Client/Server/Verify | No cumple | Cumple (nivel SIP) | 401 Security-Server + 200 Security-Verify',
    'ck/ik no llegan al UE | N/A (filtraban) | Cumple | 401 hacia UE',
    'Contact + feature-tags | No cumple | Cumple | REGISTER capturado',
    'IMPU formal (sin IMSI en URI) | No cumple | Cumple | --id sip:$MSISDN@…',
    'UAR/MAR con identidad pública (MSISDN) | N/A (403) | Cumple | logs UAR/MAR',
    'INVITE P-Preferred-Identity / P-Access-Network-Info | No cumple | Cumple | Kamailio P-CSCF MO',
    'Codecs AMR/AMR-WB (MTSI) | No cumple | Parcial | red sí; cliente no',
]
add_bullets(slide, rows, 0.8, 1.5, 12.0, 5.3, font_size=12)
add_footer(slide)

# slide 9 recommendations
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide); add_top_bar(slide)
add_title(slide, '3. Recomendaciones de corrección y seguimiento')

add_bullets(slide, [
    '✓ Configurar AKA / Digest-AKAv1-MD5 → Hecho: S-CSCF + PyHSS + sonda validan con 200 OK.',
    '✓ Negociación IPSec Gm → Nivel SIP hecho; túnel ESP pendiente por necesidad de kernel con IPsec / NET_ADMIN en Docker.',
    '✓ AMR / AMR-WB en SDP → Parcial: la red lo soporta, pero hace falta un cliente MTSI con AMR para probarlo extremo a extremo.',
    '✓ P-Access-Network-Info / P-Preferred-Identity → Hecho: P-CSCF inserta los headers en REGISTER e INVITE.',
    'Resultado E2E: scripts/test_maqueta.sh → PASS 15 / FAIL 0.',
], left=0.8, top=1.6, width=12.0, height=5.5, font_size=15)
add_footer(slide)

# slide 10 reproduction evidence
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide); add_top_bar(slide)
add_title(slide, '4. Reproducción de evidencia')

add_bullets(slide, [
    '1) Sonda AKA: docker cp maqueta-ims-3/scripts/ue_aka_probe.py mims3_pyhss_hss:/tmp/ && docker exec mims3_pyhss_hss python3 /tmp/ue_aka_probe.py',
    'Resultado esperado: "RESULTADO: AUTHENTICACIÓN AKAv1-MD5 COMPLETADA (200 OK)"',
    '2) E2E de la maqueta: bash maqueta-ims-3/scripts/test_maqueta.sh',
    'Resultado esperado: PASS: 15 FAIL: 0 ESTADO: OK',
    '3) Persistencia: parche de PyHSS en maqueta-ims-3/pyhss/diameter.py; valor por MSISDN y UAR/MAR normalizados. El archivo se monta :ro en servicios pyhss_hss y pyhss_diameter.',
], left=0.8, top=1.5, width=12.0, height=5.3, font_size=14)
add_footer(slide)

# slide 11 final conclusions
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide); add_top_bar(slide)
add_title(slide, 'Conclusión final')

add_bullets(slide, [
    'La maqueta ya no presenta los fallos críticos de autenticación, registro IMS y cabeceras estándar en la capa SIP.',
    'La parte de seguridad IPsec real queda condicionada por la capacidad del entorno para operar con NET_ADMIN y configuración del kernel, no por un fallo funcional del IMS en sí.',
    'La red soporta AMR/AMR-WB y la lógica 3GPP cumple a nivel de señalización; la limitación real está en el cliente disponible para pruebas de voz.',
    'La evidencia reproducible confirma que la solución está en un estado operativo y demostrable para validación 3GPP: PASS 15/15 en pruebas E2E.',
    'La propuesta técnica queda sólida como base para continuidad y validación más rigurosa con clientes MTSI reales.',
], left=0.8, top=1.5, width=12.0, height=5.4, font_size=14)
add_footer(slide)

prs.save(OUTPUT)
print(f'Presentación creada en {OUTPUT}')
